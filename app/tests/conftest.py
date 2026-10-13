"""测试通用夹具：内存 SQLite + 依赖覆盖 + FakeLLM + fakeredis。

对应 §9（阶段 0）+ 阶段 1 §10.1 + 阶段 2 §8：**全部测试都不需要真实 LLM/Redis**。
"""

import io
import json
from types import SimpleNamespace

import fakeredis.aioredis
import pytest
import pytest_asyncio
from httpx import ASGITransport, AsyncClient
from sqlalchemy import JSON
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import StaticPool

from app.api.v1.deps import get_stream_session_factory
from app.clients.cache import CacheBackend, RedisCache, get_cache
from app.clients.embedding import get_embedding
from app.clients.llm import LLMResult, StreamChunk, get_llm
from app.clients.queue import get_arq
from app.clients.rerank import get_reranker
from app.core.config import settings
from app.db.session import get_db
from app.main import app
from app.models.base import Base


def _make_vector_sqlite_compatible() -> None:
    """把 Vector 列在测试用的 SQLite 上降级为 JSON 存储。

    ★ 这是刻意的类型适配，不是 hack（docs/项目实现文档-阶段4.md §9.1）：
      Vector(512) 是 PostgreSQL 专属类型，SQLite 不认识，直接 create_all 会报
      "CompileError: can't render element of type VECTOR"——不加这一步，133 个存量测试
      会在 create_all 阶段全崩。
      测试关注的是**检索逻辑**（RRF 融合、降级、状态过滤），不是 pgvector 本身；
      向量距离排序由 retrieval._vector_candidates 里的应用层余弦分支走（§9.1“应用层排序”）。
      ★ none_as_null=True：Python None → SQL NULL（而 JSON 默认会把 None 存成 JSON 'null'），
        否则回填脚本的 `WHERE embedding IS NULL` 选不中未回填的行。
    """
    from app.models.chunk import Chunk

    Chunk.__table__.c.embedding.type = JSON(none_as_null=True)


@pytest.fixture(autouse=True)
def _isolate_query_rewrite():
    """阶段 4：查询改写默认关闭，隔离其余测试。

    ★ 为什么：rewrite_query 会（a）多调一次 LLM、（b）把改写句拿去检索，
      两件事都会让阶段 2/3 的会话/缓存测试误红（那些测试假定一轮只调一次 LLM）。
      改写本身由 test_query_rewrite.py 直接调用 逐例验证，在那个测试里会显式打开。
    """
    original = settings.query_rewrite_enabled
    settings.query_rewrite_enabled = False
    try:
        yield
    finally:
        settings.query_rewrite_enabled = original


def parse_sse(text: str) -> list[tuple[str, dict]]:
    """把 SSE 响应体解析成 [(event, data), ...]（阶段 2 §8.2 工具函数）。

    放在 conftest 是因为 test_stream.py 与 test_cache_saturation.py 都要用。
    """
    events: list[tuple[str, dict]] = []
    for block in text.strip().split("\n\n"):
        event, data = None, None
        for line in block.splitlines():
            if line.startswith("event:"):
                event = line[6:].strip()
            elif line.startswith("data:"):
                data = json.loads(line[5:].strip())
        if event:
            events.append((event, data))
    return events


class FakeLLM:
    """可编程的假客户端：想让它回什么就回什么。"""

    def __init__(self, content: str = "这是测试回答 [1]"):
        self.content = content
        self.calls: list[dict] = []
        self.stream_calls: list[dict] = []

    async def chat(self, messages, **kwargs) -> LLMResult:
        self.calls.append({"messages": messages, **kwargs})
        return LLMResult(
            content=self.content,
            prompt_tokens=1000,
            completion_tokens=50,
            cached_tokens=900,
            latency_ms=120,
            model="deepseek-flash",
            cost_yuan=0.0003,
            peak=False,
        )

    async def chat_stream(self, messages, **kwargs):
        """阶段 2：流式假实现——把正文按 5 字切块，最后带 usage。"""
        self.stream_calls.append({"messages": messages, **kwargs})
        text = self.content
        for i in range(0, len(text), 5):
            yield StreamChunk(kind="content", text=text[i : i + 5])
        yield StreamChunk(
            kind="usage",
            result=LLMResult(
                content="",
                prompt_tokens=1000,
                completion_tokens=50,
                cached_tokens=900,
                latency_ms=150,
                model="deepseek-flash",
                cost_yuan=0.0003,
                peak=False,
            ),
        )
        yield StreamChunk(kind="done")


@pytest_asyncio.fixture
async def fake_llm() -> FakeLLM:
    return FakeLLM()


@pytest_asyncio.fixture
async def fake_redis():
    """每个测试一个独立的假 Redis，互不污染（阶段 1 §10.1）。"""
    client = fakeredis.aioredis.FakeRedis(decode_responses=True)
    try:
        yield client
    finally:
        # 兼容不同 fakeredis 版本的关闭方式
        closer = getattr(client, "aclose", None) or getattr(client, "close", None)
        if closer is not None:
            result = closer()
            if hasattr(result, "__await__"):
                await result


@pytest_asyncio.fixture
async def test_engine():
    """内存 SQLite + StaticPool：同一连接，跨 session 可见。"""
    _make_vector_sqlite_compatible()      # ★ 必须在 create_all 之前（§9.1）
    engine = create_async_engine(
        "sqlite+aiosqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        yield engine
    finally:
        await engine.dispose()


@pytest_asyncio.fixture
def session_factory(test_engine):
    return async_sessionmaker(test_engine, expire_on_commit=False, autoflush=False)


@pytest_asyncio.fixture
async def db_session(session_factory):
    """直接给测试一个可用的 AsyncSession（不走 HTTP 端点）。"""
    async with session_factory() as s:
        yield s


@pytest_asyncio.fixture
def seed_document(db_session):
    """测试库直接建文档 + 切片（可指定 status 与每片向量），绕过上传链路。

    chunks 为 [(content, embedding 或 None), ...]。SQLite 不强制外键，owner_id 可直给。
    """
    from app.models.chunk import Chunk
    from app.models.document import Document

    async def _seed(
        *,
        owner_id: int = 1,
        status: str = "ready",
        title: str = "测试文档",
        chunks: list[tuple[str, list[float] | None]],
    ) -> tuple[Document, list[Chunk]]:
        doc = Document(
            owner_id=owner_id, title=title, source_type="txt",
            raw_text="x", status=status,
        )
        db_session.add(doc)
        await db_session.flush()
        objs: list[Chunk] = []
        for i, (content, emb) in enumerate(chunks):
            c = Chunk(
                document_id=doc.id, owner_id=owner_id, seq=i,
                content=content, char_count=len(content), embedding=emb,
            )
            db_session.add(c)
            objs.append(c)
        await db_session.flush()
        return doc, objs

    return _seed


@pytest_asyncio.fixture
def cache_backend(fake_redis) -> CacheBackend:
    """测试默认走 RedisCache(fake_redis)，限流也能通过 .redis 走真命令。"""
    return RedisCache(fake_redis)


class FakeArqPool:
    """假 ARQ 连接池：只记录投递，不真正排队（阶段 3 §9.1）。

    为什么要它：`get_arq` 是 FastAPI 依赖，未初始化时会抛
    "ARQ 队列未初始化"——而测试不走 lifespan，所以必须在夹具里覆盖掉。
    漏了这一步的后果是：任何碰 `/documents` 的测试全部失败（实测踩过，38 个用例）。
    """

    def __init__(self) -> None:
        self.jobs: list[tuple[str, tuple, dict]] = []

    async def enqueue_job(self, function: str, *args, **kwargs):
        self.jobs.append((function, args, kwargs))
        return SimpleNamespace(job_id=f"fake-job-{len(self.jobs)}")

    async def aclose(self) -> None:      # 与 ArqRedis 接口保持一致
        return None


@pytest_asyncio.fixture
async def fake_arq() -> FakeArqPool:
    return FakeArqPool()


class FakeEmbedder:
    """可编程假 embedder：按文本返回预设向量，便于精确控制召回顺序（§9.3）。

    ★ 默认返回全零向量——此时向量排序是任意的（§9.4），写断言时要么给不同向量，
      要么断言“顺序稳定”而非“顺序正确”。
    """

    def __init__(self, vectors: dict[str, list[float]] | None = None, dim: int = 512):
        self.vectors = vectors or {}
        self.dim = dim
        self.query_calls: list[str] = []
        self.doc_calls: list[list[str]] = []
        self.is_ready = True

    async def embed_query(self, text: str) -> list[float]:
        self.query_calls.append(text)
        return self.vectors.get(text, [0.0] * self.dim)

    async def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.doc_calls.append(list(texts))
        return [self.vectors.get(t, [0.0] * self.dim) for t in texts]


class FakeReranker:
    """可编程假 reranker：按 (query, doc) 返回预设分数（§9.3）。

    raises=True 时模拟精排报错，验证检索层“降级为 RRF 顺序”（§9.2）。
    """

    def __init__(
        self,
        scores: dict[tuple[str, str], float] | None = None,
        *,
        raises: bool = False,
    ):
        self.scores = scores or {}
        self.raises = raises
        self.calls: list[tuple[str, list[str]]] = []
        self.is_ready = True

    async def rerank(self, query: str, documents: list[str]) -> list[float]:
        if self.raises:
            raise RuntimeError("假 reranker 模拟精排失败")
        self.calls.append((query, documents))
        return [self.scores.get((query, d), 0.0) for d in documents]


@pytest_asyncio.fixture
async def fake_embedder() -> FakeEmbedder:
    return FakeEmbedder()


@pytest_asyncio.fixture
async def fake_reranker() -> FakeReranker:
    return FakeReranker()


@pytest_asyncio.fixture
async def client(
    fake_llm: FakeLLM,
    session_factory,
    cache_backend: CacheBackend,
    fake_arq: FakeArqPool,
    fake_embedder: FakeEmbedder,
    fake_reranker: FakeReranker,
):
    async def override_db():
        async with session_factory() as s:
            try:
                yield s
                await s.commit()
            except Exception:
                await s.rollback()
                raise

    app.dependency_overrides[get_db] = override_db
    app.dependency_overrides[get_llm] = lambda: fake_llm
    app.dependency_overrides[get_cache] = lambda: cache_backend
    # ★ 阶段 2 §6.2：流式生成器用独立 session 工厂，测试里必须换到内存 SQLite，
    #   否则它会写进真实 kb.db
    app.dependency_overrides[get_stream_session_factory] = lambda: session_factory
    # ★ 阶段 3：ARQ 依赖也必须覆盖，否则测试会去连真实 Redis 并抛
    #   "ARQ 队列未初始化"（测试不走 lifespan）
    app.dependency_overrides[get_arq] = lambda: fake_arq
    # ★ 阶段 4：embedding / rerank 依赖必须覆盖为假实现，否则 chat 端点会拿到
    #   未加载的全局客户端（_model=None）而报错。测试环境不走 lifespan。
    app.dependency_overrides[get_embedding] = lambda: fake_embedder
    app.dependency_overrides[get_reranker] = lambda: fake_reranker

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://t") as ac:
        yield ac

    app.dependency_overrides.clear()


async def register_and_login(
    client: AsyncClient, username: str = "alice", password: str = "secret123"
) -> str:
    """注册 + 登录，返回 access_token；已注册则直接登录。"""
    r = await client.post(
        "/api/v1/auth/register",
        json={"username": username, "password": password},
    )
    assert r.status_code in (201, 409), r.text
    r = await client.post(
        "/api/v1/auth/token",
        data={"username": username, "password": password},
    )
    assert r.status_code == 200, r.text
    return r.json()["access_token"]


@pytest_asyncio.fixture
async def auth_headers(client: AsyncClient) -> dict[str, str]:
    token = await register_and_login(client, "alice", "secret123")
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def unique_username():
    counter = {"n": 0}

    def _gen(prefix: str = "user") -> str:
        counter["n"] += 1
        return f"{prefix}_{counter['n']}"

    return _gen


@pytest.fixture
def make_pdf():
    """现场生成 PDF 用于测试，不往仓库里塞二进制（阶段 3 §9.2）。

    text=None 时生成只有图形的“扫描件”（无文字层），用于验证静默失败陷阱。
    encrypted=True 生成加密 PDF。reportlab 只进 dev 依赖。
    """
    def _make(text: str | None, pages: int = 1, *, encrypted: bool = False) -> bytes:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        for _ in range(pages):
            if text:
                # 不依赖中文字体：ASCII 文字层同样能被 pypdf 提取，足以验证行为
                c.setFont("Helvetica", 14)
                y = 780
                for line in text.split("\n"):
                    c.drawString(60, y, line)
                    y -= 26
            else:
                c.rect(100, 500, 300, 200, fill=1)   # 只有图形，无文字层
            c.showPage()
        c.save()
        data = buf.getvalue()
        if encrypted:
            from pypdf import PdfReader, PdfWriter

            writer = PdfWriter()
            reader = PdfReader(io.BytesIO(data))
            writer.append(reader)
            writer.encrypt("secret-password")
            out = io.BytesIO()
            writer.write(out)
            return out.getvalue()
        return data

    return _make


@pytest.fixture
def make_docx():
    """现场生成含段落 + 表格的 .docx（验证表格内容也被提取）。"""
    def _make(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
        import docx

        document = docx.Document()
        for p in paragraphs:
            document.add_paragraph(p)
        if table_rows:
            table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
            for r, row in enumerate(table_rows):
                for c, val in enumerate(row):
                    table.cell(r, c).text = val
        buf = io.BytesIO()
        document.save(buf)
        return buf.getvalue()

    return _make
