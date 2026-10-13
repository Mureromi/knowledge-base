"""阶段 3：文档文本提取。txt/md/pdf/docx 统一入口。

★ 所有 PDF/DOCX 解析都是**同步阻塞**的 CPU 活，调用方（worker）必须用
  asyncio.to_thread 包装（见 §1）。在 async 事件循环里直接跑会卡死整个 worker 进程，
  而且这个 bug 测试测不出来——只有并发冒烟测试（§9.4）能兜住。

两类确定性失败（§2.2 / §2.3）：
- EmptyDocumentError：扫描件/图片型 PDF——"提取成功但内容为空"，重试也没用
- UnsupportedDocumentError：加密、损坏、超页数——同样不该重试
未知异常留给调用方判定是否可重试。
"""

import io
import logging

from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.services.ingest import decode_bytes

logger = logging.getLogger(__name__)


class EmptyDocumentError(Exception):
    """没有可提取的文字（扫描件/图片型 PDF）。确定性失败，不重试。"""


class UnsupportedDocumentError(Exception):
    """加密、损坏、超出页数限制等。确定性失败，不重试。"""


def extract_pdf_text(raw: bytes, *, max_pages: int) -> str:
    """从 PDF 提取文字。**同步阻塞函数**，调用方负责 to_thread。"""
    try:
        reader = PdfReader(io.BytesIO(raw))
    except PdfReadError as exc:
        raise UnsupportedDocumentError(f"PDF 无法打开：{exc}") from exc
    except Exception as exc:
        raise UnsupportedDocumentError(f"PDF 无法打开：{type(exc).__name__}") from exc

    if reader.is_encrypted:
        # 尝试空密码（有些 PDF 只是限制编辑，允许读取）
        try:
            if reader.decrypt("") == 0:
                raise UnsupportedDocumentError("PDF 已加密，需要密码才能读取")
        except UnsupportedDocumentError:
            raise
        except Exception as exc:
            raise UnsupportedDocumentError("PDF 已加密，需要密码才能读取") from exc

    page_count = len(reader.pages)
    if page_count > max_pages:
        raise UnsupportedDocumentError(
            f"PDF 页数 {page_count} 超过上限 {max_pages} 页"
        )

    parts: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001
            # 单页失败不放弃整份文档：真实 PDF 常有坏页（图表页、嵌入对象）
            logger.warning("第 %s 页提取失败：%s", index + 1, exc)
            parts.append("")

    text = "\n\n".join(p.strip() for p in parts if p.strip())

    # ★ 关键检查（§2.2）：扫描件会"成功返回空字符串"，必须显式判失败，
    #   否则以 status=ready、chunk_count=0 入库 → 用户看不到任何错误、检索永远落空
    if not text.strip():
        raise EmptyDocumentError(
            "该 PDF 没有可提取的文字层，可能是扫描件或图片型 PDF。"
            "本项目暂不支持 OCR，请上传含文字层的 PDF，或先转成 txt/md。"
        )
    return text


def extract_docx_text(raw: bytes) -> str:
    """从 .docx 提取文字。段落 + 表格（政策/价格类文档大量信息在表格里）。"""
    import docx  # 延迟导入，减少不需要 docx 时的 worker 启动开销

    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as exc:
        raise UnsupportedDocumentError(f"docx 无法打开：{exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    if not text.strip():
        raise EmptyDocumentError("docx 中没有可提取的文字")
    return text


def extract_text(raw: bytes, source_type: str, max_pages: int) -> str:
    """按 source_type 分派的统一入口。**同步阻塞**，worker 侧用 to_thread 调用。

    对 txt/md 也做兼容（虽然端点走快路径一般不会到这里），保证 worker
    重新解析历史任务时四种类型都能处理。
    """
    if source_type == "pdf":
        return extract_pdf_text(raw, max_pages=max_pages)
    if source_type == "docx":
        return extract_docx_text(raw)
    if source_type in ("txt", "md"):
        text = decode_bytes(raw)
        if not text.strip():
            raise EmptyDocumentError("文本文件内容为空")
        return text
    raise UnsupportedDocumentError(f"不支持的文档类型：{source_type}")
